# -*- coding: utf-8 -*-
"""
trusscore_docx.py - canonical Trusscore Word document builder.

Implements references/docx.md in code so the docx layer is never re-derived
by hand: Aptos everywhere (including the List Bullet / List Number styles,
which do NOT inherit the Normal override), TC Slate headings, TC Text body,
0.8"/0.9" margins, slate-underlined table headers with grey row dividers,
yellow-tint callouts, and a deterministic qa_verify() post-build check.

Usage:
    import trusscore_docx as tdx
    doc = tdx.new_doc()
    tdx.title_block(doc, "Document Title", ["Date: ...", "From: ..."])
    tdx.heading(doc, "Section Title", 1)
    tdx.para(doc, "Body copy.", bold_lead="Label.")
    tdx.bullet(doc, "A fragment bullet, no terminal period")
    tdx.numbered(doc, "First step")
    tdx.table(doc, ["Col A", "Col B"], [["a", "b"], ["c", "d"]])
    tdx.callout(doc, "One important takeaway sentence.")
    tdx.divider(doc)
    tdx.save(doc, "out.docx")
    ok, issues = tdx.qa_verify("out.docx")
"""
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Aptos"
SLATE = RGBColor(0x3A, 0x4B, 0x5C)
TEXT = RGBColor(0x11, 0x1A, 0x22)
MUTED = RGBColor(0x6A, 0x7A, 0x88)
GREY_HEX = "DFDBDA"          # dividers / table row borders
TINT_HEX = "FFF0C2"          # TC Yellow 20% callout fill
EM = "—"

PATCH_STYLES = ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3",
                "List Bullet", "List Bullet 2", "List Number", "List Number 2"]
HEADINGS = {"Title": (20, True), "Heading 1": (16, True),
            "Heading 2": (13, True), "Heading 3": (11.5, True)}


def _force_font(style, name=FONT):
    """Set the font on a style at the XML level so nothing falls back to the theme."""
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def new_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.9)
    for name in PATCH_STYLES:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _force_font(style)
        if name in HEADINGS:
            size, bold = HEADINGS[name]
            style.font.size = Pt(size)
            style.font.bold = bold
            style.font.color.rgb = SLATE
            pf = style.paragraph_format
            pf.space_before = Pt(12 if name == "Heading 1" else 10)
            pf.space_after = Pt(4)
        else:
            style.font.size = Pt(11)
            style.font.color.rgb = TEXT
            pf = style.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            pf.line_spacing = 1.25
    return doc


def save(doc, path):
    doc.save(path)
    return path


def heading(doc, text, level=1):
    return doc.add_paragraph(text, style="Heading %d" % level if level else "Title")


def title_block(doc, title, meta_lines=()):
    doc.add_paragraph(title, style="Title")
    for line in meta_lines:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        p.paragraph_format.space_after = Pt(2)
    divider(doc)
    return doc


def para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    p.add_run(text)
    return p


def bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    return doc.add_paragraph(str(text).strip().rstrip("."), style=style)


def numbered(doc, text, level=0):
    style = "List Number" if level == 0 else "List Number 2"
    return doc.add_paragraph(text, style=style)


def _cell_border(cell, edge, hex_color, size_8ths):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = tc_pr.makeelement(qn("w:tcBorders"), {})
        tc_pr.append(borders)
    el = borders.find(qn("w:" + edge))
    if el is None:
        el = borders.makeelement(qn("w:" + edge), {})
        borders.append(el)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size_8ths))
    el.set(qn("w:color"), hex_color)


def _shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _cell_text(cell, text, bold=False, color=TEXT, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color


def table(doc, headers, rows):
    """Spec table: slate 2pt header underline, 1pt grey row dividers, no verticals."""
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    t.style = doc.styles["Normal Table"]      # strip default grid styling
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        _cell_text(cell, h, bold=True, color=SLATE)
        _cell_border(cell, "bottom", "3A4B5C", 16)   # 2pt slate header underline
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            cell = t.rows[i].cells[j]
            _cell_text(cell, val)
            _cell_border(cell, "bottom", GREY_HEX, 8)   # 1pt grey divider
    return t


def callout(doc, text):
    """Yellow-tint callout box (full sentence: keeps terminal punctuation)."""
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles["Normal Table"]
    cell = t.rows[0].cells[0]
    _shade(cell, TINT_HEX)
    for edge in ("top", "bottom", "left", "right"):
        _cell_border(cell, edge, TINT_HEX, 4)
    _cell_text(cell, text, bold=False, color=SLATE, size=10.5)
    return t


def classification(doc, label="INTERNAL"):
    """Classification label: TC Slate, caps, letter-spaced - never the yellow eyebrow."""
    p = doc.add_paragraph()
    r = p.add_run(" ".join(label.upper()))
    r.font.name = FONT
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = SLATE
    return p


def divider(doc):
    """1pt TC Grey horizontal rule."""
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    borders = p_pr.makeelement(qn("w:pBdr"), {})
    bottom = p_pr.makeelement(qn("w:bottom"), {})
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), GREY_HEX)
    borders.append(bottom)
    p_pr.append(borders)
    return p


# ---------------- QA ----------------

def qa_verify(path):
    """Deterministic post-build check per references/docx.md. Returns (ok, issues)."""
    issues = []
    doc = Document(path)

    # 1. Styles carry Aptos and slate headings
    for name in PATCH_STYLES:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        if style.font.name != FONT:
            issues.append("style '%s' font is %r, not %s" % (name, style.font.name, FONT))
        if name in HEADINGS and style.font.color.rgb != SLATE:
            issues.append("style '%s' color is %s, not 3A4B5C" % (name, style.font.color.rgb))

    # 2. No run overrides an off-spec font
    for p in doc.paragraphs:
        for r in p.runs:
            if r.font.name is not None and not r.font.name.startswith(FONT):
                issues.append("run font %r in %r" % (r.font.name, p.text[:40]))

    # 3. Em dashes: max two, never spaced
    full = "\n".join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                full += "\n" + c.text
    n_em = full.count(EM)
    if n_em > 2:
        issues.append("%d em dashes (max two)" % n_em)
    if re.search(r"[ \t]" + EM + r"|" + EM + r"[ \t]", full):
        issues.append("spaced em dash present")

    # 4. Tables have the spec borders
    for ti, t in enumerate(doc.tables):
        hdr = t.rows[0].cells[0]._tc
        if hdr.find(qn("w:tcPr")) is None or \
           hdr.find(qn("w:tcPr")).find(qn("w:tcBorders")) is None:
            issues.append("table %d header cell has no borders" % ti)

    # 5. Bullets: no terminal period
    for p in doc.paragraphs:
        if p.style.name.startswith("List Bullet") and p.text.strip().endswith("."):
            issues.append("bullet ends with period: %r" % p.text[:40])

    return (not issues), issues


if __name__ == "__main__":
    import sys
    ok, issues = qa_verify(sys.argv[1])
    print("QA PASS" if ok else "QA FAIL")
    for i in issues:
        print(" -", i)
    sys.exit(0 if ok else 1)
